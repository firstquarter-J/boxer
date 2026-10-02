"""채팅과 같은 조회 스냅샷의 전체 분석 행을 수식 없는 엑셀로 내보낸다."""

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from boxer_company.assistant.contracts import AssistantFile
from boxer_company.assistant.file_contracts import MAX_REPORT_FILE_BYTES

_SHEETS = (
    ("병원별 녹화", "hospitals", "totalCount", "건"),
    ("병원별 신규 바코드", "hospitals", "newBarcodeCount", "개"),
    ("병실별 녹화", "rooms", "totalCount", "건"),
    ("병실별 신규 바코드", "rooms", "newBarcodeCount", "개"),
)
_HEADER_FILL = PatternFill("solid", fgColor="203864")


class ReportFileTooLargeError(ValueError):
    """전체 행을 조용히 자르지 않고 조회 범위를 줄이도록 안내한다."""


def _append(sheet: Any, values: list[Any]) -> None:
    # openpyxl이 셀을 만들 때 제어 문자를 검사하므로 값 바인딩 전에 제거한다.
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # DB 이름이 '='·URL로 시작해도 수식·외부 링크로 실행하지 않는다.
            cell.data_type = "s"


def _header(sheet: Any, row: int) -> None:
    for cell in sheet[row]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 42


def build_recordings_trend_excel(summary: dict[str, Any], *, now: datetime) -> AssistantFile:
    workbook = Workbook()
    info = workbook.active
    info.title = "조회 조건"
    complete = [week for week in summary["weeks"] if week["complete"]]
    required = summary["declineWeeks"]
    metadata = [
        ["항목", "내용"],
        ["조회 기간 (KST)", f"{summary['startDate']} ~ {summary['endDate']}"],
        ["조회 시각 (KST)", now.strftime("%Y-%m-%d %H:%M:%S")],
        ["대상 병원", ", ".join(summary["hospitalNames"]) or "전체"],
        ["연속 감소 조건", f"마지막 완료 주까지 {required}주 이상" if required else "없음 (증가·유지 포함)"],
        ["매주 감소 기준", f"{summary['dropPercent']:g}% 이상, {summary['minimumDrop']}건/개 이상 (보합 제외)"],
        ["집계 기준", "채팅 응답과 동일한 조회 결과. 지정 조건에 맞는 모든 행 포함."],
        ["신규 바코드", "전체 촬영 이력의 최초 1회를 최초 촬영 병원·병실에만 집계"],
        ["부분 주", "부분 주·진행 중인 주는 수량만 표시하고 증감·연속 감소 판정에서 제외"],
        ["병실 미지정", "병원 합계에만 포함. 아래 주별 합계에서 미지정 수량을 확인할 수 있음"],
        ["증감률 빈칸", "비교할 완료 주가 부족하거나 기준 주 수량이 0이라 비율 계산 불가"],
        ["정렬", "첫 완료 주 대비 변화량의 절댓값 내림차순"],
    ]
    if len(complete) >= 2:
        for label, first in (("전체 변화 비교", complete[0]), ("최근 전주 비교", complete[-2])):
            metadata.append([label, (f"{first['startDate']}~{first['endDate']} → "
                                    f"{complete[-1]['startDate']}~{complete[-1]['endDate']}")])
    for row in metadata:
        _append(info, row)
    _header(info, 1)
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 105
    info.freeze_panes = "B2"
    _append(info, [])
    _append(info, ["시트", "조건 적용 대상 수", "기간 내 촬영 대상 수"])
    _header(info, info.max_row)

    for title, group, metric, unit in _SHEETS:
        all_rows = summary["entityTrends"][metric][group]
        rows = [row for row in all_rows if row["declineStreak"] >= required] if required else all_rows
        _append(info, [title, len(rows), len(all_rows)])
        sheet = workbook.create_sheet(title)
        headers = ["병원 ID", "병실 ID", "대상 이름"]
        headers += [f"{week['startDate']}~{week['endDate']} ({unit})" +
                    (" [부분 주]" if not week["complete"] else "") for week in summary["weeks"]]
        headers += ["흐름", f"첫 완료 주 대비 증감 ({unit})", "첫 완료 주 대비 증감률",
                    f"최근 전주 대비 증감 ({unit})", "최근 전주 대비 증감률", "조건 충족 연속 감소 (주)"]
        _append(sheet, headers)
        _header(sheet, 1)
        for row in rows:
            key = row["key"]
            hospital, room = key if group == "rooms" else (key, None)
            _append(sheet, [hospital, room, row["name"], *row["counts"], row["direction"], row["netDelta"],
                            row["netRate"] / 100 if row["netRate"] is not None else None,
                            row["latestDelta"], row["latestRate"] / 100 if row["latestRate"] is not None else None,
                            row["declineStreak"]])
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = sheet.column_dimensions["B"].width = 14
        sheet.column_dimensions["C"].width = 45
        for index in range(4, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 23
        for cells in sheet.iter_rows(min_row=2, min_col=4):
            for cell in cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0%" if "증감률" in headers[cell.column - 1] else "#,##0"

    # 조건 적용 전 전체 합계를 별도로 남겨 병실 미지정·부분 주를 검산할 수 있게 한다.
    _append(info, [])
    _append(info, ["주 시작", "주 종료", "비교 포함", "녹화", "신규 바코드", "병실 미지정 녹화", "병실 미지정 신규 바코드"])
    _header(info, info.max_row)
    for week in summary["weeks"]:
        _append(info, [week["startDate"], week["endDate"], "완료 주" if week["complete"] else "부분 주",
                       week["totalCount"], week["newBarcodeCount"],
                       week["totalCount"] - sum(row["totalCount"] for row in week["rooms"].values()),
                       week["newBarcodeCount"] - sum(row["newBarcodeCount"] for row in week["rooms"].values())])
    # 임시 파일이나 공유 URL 없이 요청 수명 안에서만 생성·전달한다.
    with BytesIO() as output:
        workbook.save(output)
        content = output.getvalue()
    workbook.close()
    if len(content) > MAX_REPORT_FILE_BYTES:
        raise ReportFileTooLargeError("report file exceeds size limit")
    return AssistantFile(filename=f"recordings-trend-{summary['startDate']}_{summary['endDate']}.xlsx", content=content)
