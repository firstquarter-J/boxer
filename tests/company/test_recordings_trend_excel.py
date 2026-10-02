"""동일 조회 스냅샷의 엑셀 전체 행, 안전한 파일 계약과 requester DM 전달을 검증한다."""

import json
import logging
from copy import deepcopy
from dataclasses import replace
from io import BytesIO
from unittest.mock import Mock, patch

import pytest
from openpyxl import load_workbook

from boxer_company._operation_routing_file import match_device_file_operation_route
from boxer_company.assistant.contracts import (
    AssistantFile,
    AssistantMessage,
    CompanyAssistantResult,
)
from boxer_company.assistant.file_contracts import (
    MAX_FILE_RESPONSE_BYTES,
    MAX_REPORT_FILE_BYTES,
)
from boxer_company.assistant.operational_read_routes import (
    WeeklyRecordingsSummaryAssistantRoute,
)
from boxer_company.assistant.recordings_trend_format import format_recordings_trend
from boxer_company.read_routing import match_weekly_recordings_summary_route
from boxer_company.recordings_trend_excel import (
    ReportFileTooLargeError,
    build_recordings_trend_excel,
)
from boxer_company_adapter_slack.assistant_bridge import render_company_assistant_result
from boxer_company_adapter_slack.company_api_client import (
    CompanyApiContractError,
    _deserialize_result,
    _deserialize_result_payload,
)
from boxer_company_api.schemas import serialize_result
from tests.company import test_recordings_trend as trend_tests

breakdown_summary = trend_tests.breakdown_summary
recordings_db = trend_tests.recordings_db
NOW = trend_tests.NOW


def _result(file):
    return CompanyAssistantResult(route="weekly_recordings_summary", outcome="answered",
                                  messages=(AssistantMessage("분석 결과"),), files=(file,))


def test_workbook_has_four_complete_series_and_numeric_percentages(breakdown_summary):
    file = build_recordings_trend_excel(breakdown_summary, now=NOW)
    workbook = load_workbook(BytesIO(file.content))
    assert workbook.sheetnames == ["조회 조건", "병원별 녹화", "병원별 신규 바코드", "병실별 녹화", "병실별 신규 바코드"]
    hospital = workbook["병원별 신규 바코드"]
    rows = {row[0]: row for row in hospital.iter_rows(min_row=2, values_only=True)}
    assert rows[1][3:7] == (40, 30, 20, 10)
    assert rows[1][8:13] == (-30, -0.75, -10, -0.5, 3)
    assert rows[3][3:7] == (0, 0, 0, 0)
    rooms = workbook["병실별 신규 바코드"]
    room_rows = {(row[0], row[1]): row for row in rooms.iter_rows(min_row=2, values_only=True)}
    assert room_rows[(1, 11)][3:7] == (30, 20, 10, 0)
    assert room_rows[(2, 21)][3:7] == (2, 4, 6, 8)
    assert rooms.freeze_panes == "D2" and rooms.auto_filter.ref == rooms.dimensions
    assert hospital["J2"].number_format == "0.0%"
    assert workbook["조회 조건"]["B2"].value == "2026-08-24 ~ 2026-09-20"
    assert len(file.content) < MAX_REPORT_FILE_BYTES
    workbook.close()


def test_decline_filter_is_applied_independently_in_every_sheet(breakdown_summary):
    breakdown_summary["declineWeeks"] = 2
    workbook = load_workbook(BytesIO(build_recordings_trend_excel(breakdown_summary, now=NOW).content))
    expected = [(3, None), (1, None), (3, 31), (1, 11)]
    for sheet, identity in zip(workbook.worksheets[1:], expected, strict=True):
        assert sheet.max_row == 2
        assert (sheet["A2"].value, sheet["B2"].value) == identity
    workbook.close()


def test_export_keeps_rows_omitted_from_chat(breakdown_summary):
    # 실제 운영 규모보다 큰 목록을 만들어 채팅 예산과 파일 행 수가 독립인지 확인한다.
    template = breakdown_summary["entityTrends"]["totalCount"]["rooms"][0]
    rows = [{**template, "key": (1, index + 1), "name": f"병원 · {index + 1}병실"} for index in range(2000)]
    breakdown_summary["entityTrends"]["totalCount"]["rooms"] = rows
    assert "응답 길이 제한" in format_recordings_trend(breakdown_summary, now=NOW)[2]
    workbook = load_workbook(BytesIO(build_recordings_trend_excel(breakdown_summary, now=NOW).content))
    assert workbook["병실별 녹화"].max_row == 2001
    assert workbook["병실별 녹화"]["B2001"].value == 2000
    workbook.close()


def test_names_are_literal_strings_and_partial_weeks_are_marked(breakdown_summary):
    name = '=HYPERLINK("https://invalid.example", "외부 링크")'
    breakdown_summary["hospitalNames"] = [name]
    for row in breakdown_summary["entityTrends"]["totalCount"]["hospitals"]:
        row["name"] = name + "\x01"
    breakdown_summary["weeks"][0]["complete"] = False
    workbook = load_workbook(BytesIO(build_recordings_trend_excel(breakdown_summary, now=NOW).content))
    assert workbook["병원별 녹화"]["C2"].value == name
    assert workbook["병원별 녹화"]["C2"].data_type == "s"
    assert "[부분 주]" in workbook["병원별 녹화"]["D1"].value
    assert all(cell.data_type != "f" and cell.hyperlink is None for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_route_attaches_same_single_query_snapshot(recordings_db):
    with patch("boxer_company.assistant.operational_read_routes._coerce_weekly_recordings_report_now", return_value=NOW):
        result = WeeklyRecordingsSummaryAssistantRoute().handle(trend_tests.request("최근 4주 녹화 추이"))
    assert len(recordings_db) == 1 and len(result.files) == 1
    # API → Slack 역직렬화 후에도 원본 파일 바이트와 네 분석 본문을 그대로 보존한다.
    restored = _deserialize_result_payload(serialize_result(result, "excel-test"), "excel-test")
    assert restored.files == result.files
    assert len(restored.messages) == 4
    workbook = load_workbook(BytesIO(restored.files[0].content))
    assert workbook["병원별 녹화"].max_row > 1
    workbook.close()


@pytest.mark.parametrize("failure", [RuntimeError("private-provider-detail"), ReportFileTooLargeError("too large")])
def test_export_failure_preserves_analysis_without_requery(recordings_db, failure):
    with patch("boxer_company.assistant.operational_read_routes._coerce_weekly_recordings_report_now", return_value=NOW), patch(
        "boxer_company.assistant.operational_read_routes.build_recordings_trend_excel", side_effect=failure,
    ):
        result = WeeklyRecordingsSummaryAssistantRoute().handle(trend_tests.request("최근 4주 녹화 추이"))
    assert result.outcome == "answered" and len(result.messages) == 4 and not result.files
    assert "엑셀 파일" in result.messages[0].body and "private-provider-detail" not in result.messages[0].body
    assert len(recordings_db) == 1


def test_no_data_still_exports_headers_and_query_conditions(recordings_db):
    with patch("boxer_company.assistant.operational_read_routes._coerce_weekly_recordings_report_now", return_value=NOW):
        result = WeeklyRecordingsSummaryAssistantRoute().handle(trend_tests.request("2026-07-01 ~ 2026-07-31 녹화 추이"))
    assert result.outcome == "no_evidence" and len(result.files) == 1
    workbook = load_workbook(BytesIO(result.files[0].content))
    assert all(sheet.max_row == 1 for sheet in workbook.worksheets[1:])
    workbook.close()


@pytest.mark.parametrize("question", [
    "최근 4주 녹화 추이 엑셀 다운로드", "8월부터 지금까지 추이 엑셀로 줘",
    "2주 이상 감소 데이터 조회 xlsx", "병원: A병원 최근 4주 녹화 추이 Excel download",
])
def test_excel_phrases_do_not_route_to_device_download(question):
    request = trend_tests.request(question)
    assert match_weekly_recordings_summary_route(request) == "weekly_recordings_summary"
    assert match_device_file_operation_route(replace(request, metadata={"route_group": "operations"})) is None


@pytest.mark.parametrize("question", [
    "최근 4주 녹화 다운로드 추이", "최근 4주 녹화 추이 엑셀 다운로드 영상 다운로드",
    "12345678910 최근 4주 녹화 추이 엑셀 다운로드", "장비 파일 다운로드",
])
def test_excel_support_does_not_capture_video_download(question):
    assert match_weekly_recordings_summary_route(trend_tests.request(question)) is None


@pytest.mark.parametrize(("key", "bad_value"), [
    ("filename", "../report.xlsx"), ("deliveryScope", "conversation"), ("mediaType", "text/html"),
    ("contentBase64", "not base64"), ("sha256", "0" * 64), ("sizeBytes", True), ("sizeBytes", 4),
])
def test_client_rejects_invalid_file_manifest(breakdown_summary, key, bad_value):
    payload = serialize_result(_result(build_recordings_trend_excel(breakdown_summary, now=NOW)), "test")
    payload["files"][0][key] = bad_value
    with pytest.raises(CompanyApiContractError):
        _deserialize_result_payload(payload, "test")


def test_file_bytes_have_separate_budget_and_are_not_logged_in_repr(breakdown_summary):
    # 압축률과 무관하게 전송 최대치에 가까운 bytes로 메시지 예산을 검사한다.
    file = AssistantFile("recordings-trend-2026-08-24_2026-09-20.xlsx", b"PK\x03\x04" + b"x" * (MAX_REPORT_FILE_BYTES - 4))
    result = replace(_result(file), messages=tuple(AssistantMessage(body) for body in format_recordings_trend(breakdown_summary, now=NOW)))
    payload = serialize_result(result, "test")
    encoded = json.dumps(payload, ensure_ascii=False).encode()
    assert 1_048_576 < len(encoded) < MAX_FILE_RESPONSE_BYTES
    response = Mock(headers={"content-type": "application/json"}, content=encoded)
    response.json.return_value = payload
    restored = _deserialize_result(response, "test")
    assert restored.files == (file,) and len(restored.messages) == 4
    assert "content=" not in repr(restored)
    without_file = deepcopy(payload)
    without_file.pop("files")
    response.json.return_value = without_file
    with pytest.raises(CompanyApiContractError):
        _deserialize_result(response, "test")
    with pytest.raises(ValueError):
        serialize_result(replace(result, files=(replace(file, content=file.content + b"x"),)), "test")


def test_file_scope_route_and_count_are_closed(breakdown_summary):
    result = _result(build_recordings_trend_excel(breakdown_summary, now=NOW))
    for bad_result in (replace(result, route="device_detail"), replace(result, files=result.files * 2)):
        with pytest.raises(ValueError):
            serialize_result(bad_result, "test")
    payload = serialize_result(result, "test")
    for bad_files in ([], payload["files"] * 2, [{**payload["files"][0], "url": "https://invalid.example"}]):
        with pytest.raises(CompanyApiContractError):
            _deserialize_result_payload({**payload, "files": bad_files}, "test")
    with pytest.raises(CompanyApiContractError):
        _deserialize_result_payload({**payload, "route": "device_detail"}, "test")


def test_renderer_uploads_exact_bytes_only_to_requester_dm(breakdown_summary):
    file = build_recordings_trend_excel(breakdown_summary, now=NOW)
    client, reply = Mock(), Mock()
    client.conversations_open.return_value = {"ok": True, "channel": {"id": "D123"}}
    client.files_upload_v2.return_value = {"ok": True}
    render_company_assistant_result(_result(file), reply=reply, actor_id="U123", client=client, logger=logging.getLogger(__name__))
    client.conversations_open.assert_called_once_with(users=["U123"])
    client.files_upload_v2.assert_called_once()
    kwargs = client.files_upload_v2.call_args.kwargs
    assert kwargs["channel"] == "D123" and kwargs["file"] == file.content and kwargs["filename"] == file.filename
    assert "DM으로 보냈어" in reply.call_args.args[0]
    client.chat_postMessage.assert_not_called()


@pytest.mark.parametrize("failure", ["upload", "dm", "actor", "channel"])
def test_upload_failure_never_falls_back_to_public_or_retries(breakdown_summary, failure, caplog):
    file = build_recordings_trend_excel(breakdown_summary, now=NOW)
    client, reply = Mock(), Mock()
    client.conversations_open.return_value = {"ok": True, "channel": {"id": "C123" if failure == "channel" else "D123"}}
    client.files_upload_v2.side_effect = RuntimeError("secret-file-content")
    if failure == "dm":
        client.conversations_open.side_effect = RuntimeError("private-dm-error")
    render_company_assistant_result(_result(file), reply=reply, actor_id=None if failure == "actor" else "U123",
                                    client=client, logger=logging.getLogger(__name__))
    assert "보내지 못했어" in reply.call_args.args[0]
    assert client.files_upload_v2.call_count == (1 if failure == "upload" else 0)
    assert "secret-file-content" not in caplog.text and "private-dm-error" not in caplog.text
